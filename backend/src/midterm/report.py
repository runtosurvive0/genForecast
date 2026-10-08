"""연간 결과 → 당진본부 일·월 발전량 표, 엑셀, 화면용 JSON."""

from __future__ import annotations

import json
from collections import defaultdict
from datetime import date
from pathlib import Path

import numpy as np

from midterm.annual import AnnualResult
from midterm.fleet import DANGJIN
from midterm.units import plant_of


def _dangjin_index(result: AnnualResult) -> list[int]:
    order = sorted((j for j, n in enumerate(result.names) if plant_of(n) == DANGJIN),
                   key=lambda j: int(result.names[j][len(DANGJIN):]))
    return order


def summarize(result: AnnualResult) -> dict:
    """화면·엑셀이 같이 쓰는 표. MW 는 시간평균, MWh 는 합계."""
    dj = _dangjin_index(result)
    names = [result.names[j] for j in dj]
    capacity = {"당진1": 500, "당진2": 500, "당진3": 500, "당진4": 500, "당진5": 500,
                "당진6": 500, "당진7": 500, "당진8": 500, "당진9": 1020, "당진10": 1020}
    installed = sum(capacity.get(n, 0) for n in names)
    out = result.output[:, dj]
    on = result.online[:, dj]
    st = result.starts[:, dj]
    total_coal = result.output.sum(axis=1)

    daily = []
    for i, day in enumerate(result.days):
        rows = slice(24 * i, 24 * i + 24)
        mwh = float(out[rows].sum())
        daily.append({
            "date": day.isoformat(),
            "dangjin_mwh": round(mwh, 1),
            "dangjin_avg_mw": round(mwh / 24, 1),
            "dangjin_cf_pct": round(mwh / (installed * 24) * 100, 2),
            "units_online_avg": round(float(on[rows].sum(axis=1).mean()), 2),
            "starts": int(st[rows].sum()),
            "unit_mwh": {n: round(float(out[rows, k].sum()), 1) for k, n in enumerate(names)},
            "net_demand_avg_mw": round(float(result.net[rows].mean()), 0),
            "nuclear_avg_mw": round(float(result.nuclear[rows].mean()), 0),
            "coal_target_avg_mw": round(float(result.coal_target[rows].mean()), 0),
            "coal_solved_avg_mw": round(float(total_coal[rows].mean()), 0),
            "mismatch_mwh": round(float((result.shortage[rows] + result.excess[rows]).sum()), 1),
        })

    by_month: dict[str, list[int]] = defaultdict(list)
    for i, day in enumerate(result.days):
        by_month[day.strftime("%Y-%m")].append(i)
    monthly = []
    for month, idx in sorted(by_month.items()):
        hours = np.concatenate([np.arange(24 * i, 24 * i + 24) for i in idx])
        mwh = float(out[hours].sum())
        monthly.append({
            "month": month, "days": len(idx),
            "dangjin_mwh": round(mwh, 0),
            "dangjin_avg_mw": round(mwh / len(hours), 1),
            "dangjin_cf_pct": round(mwh / (installed * len(hours)) * 100, 2),
            "daily_mwh_min": round(min(daily[i]["dangjin_mwh"] for i in idx), 0),
            "daily_mwh_max": round(max(daily[i]["dangjin_mwh"] for i in idx), 0),
            "units_online_avg": round(float(on[hours].sum(axis=1).mean()), 2),
            "starts": int(st[hours].sum()),
            "unit_avg_mw": {n: round(float(out[hours, k].mean()), 1) for k, n in enumerate(names)},
            "unit_cf_pct": {n: round(float(out[hours, k].mean()) / capacity.get(n, 500) * 100, 1)
                            for k, n in enumerate(names)},
            "unit_hours_online": {n: int(on[hours, k].sum()) for k, n in enumerate(names)},
            "demand_avg_mw": round(float(result.demand[hours].mean()), 0),
            "solar_avg_mw": round(float(result.solar[hours].mean()), 0),
            "net_demand_avg_mw": round(float(result.net[hours].mean()), 0),
            "nuclear_avg_mw": round(float(result.nuclear[hours].mean()), 0),
            "coal_model_avg_mw": round(float(result.coal_model[hours].mean()), 0),
            "coal_target_avg_mw": round(float(result.coal_target[hours].mean()), 0),
            "coal_solved_avg_mw": round(float(total_coal[hours].mean()), 0),
            "dangjin_share_pct": round(mwh / max(float(total_coal[hours].sum()), 1) * 100, 2),
            "mismatch_mwh": round(float((result.shortage[hours] + result.excess[hours]).sum()), 0),
        })

    hours_total = len(result.days) * 24
    total_mwh = float(out.sum())
    statuses = defaultdict(int)
    for s in result.window_status:
        statuses[s] += 1
    return {
        "scenario": result.scenario.to_json(),
        "units": names, "installed_mw": installed,
        "annual": {
            "dangjin_mwh": round(total_mwh, 0),
            "dangjin_avg_mw": round(total_mwh / hours_total, 1),
            "dangjin_cf_pct": round(total_mwh / (installed * hours_total) * 100, 2),
            "starts": int(st.sum()),
            "coal_solved_twh": round(float(total_coal.sum()) / 1e6, 2),
            "coal_target_twh": round(float(result.coal_target.sum()) / 1e6, 2),
            "mismatch_gwh": round(float((result.shortage + result.excess).sum()) / 1e3, 1),
            "net_demand_avg_mw": round(float(result.net.mean()), 0),
            "nuclear_avg_mw": round(float(result.nuclear.mean()), 0),
            "dangjin_fuel_cost_billion_won": round(float(result.fuel_cost_won[:, dj].sum()) / 1e9, 1),
        },
        "monthly": monthly, "daily": daily,
        "solve": {"elapsed_seconds": round(result.elapsed_seconds, 1),
                  "windows": len(result.window_status), "window_status": dict(statuses)},
        "model_info": result.model_info, "warnings": result.warnings,
        "outages": {
            "nuclear": [o.as_row() for o in result.scenario.nuclear_oh],
            "coal": [o.as_row() for o in result.scenario.coal_oh]},
    }


def write_excel(summary: dict, path: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(color="FFFFFF", bold=True)

    def sheet(title, header, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "B2"
        for i, _ in enumerate(header, start=1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 12)
        return ws

    wb.remove(wb.active)
    units = summary["units"]
    sheet("월별", ["월", "일수", "당진 발전량(MWh)", "당진 평균출력(MW)", "이용률(%)",
                  "일발전량 최소(MWh)", "일발전량 최대(MWh)", "평균 운전대수", "기동횟수",
                  "당진/전국석탄(%)", "순수요(MW)", "원전(MW)", "석탄목표(MW)", "석탄해(MW)",
                  "목표불일치(MWh)"] + [f"{u} 평균(MW)" for u in units],
          [[m["month"], m["days"], m["dangjin_mwh"], m["dangjin_avg_mw"], m["dangjin_cf_pct"],
            m["daily_mwh_min"], m["daily_mwh_max"], m["units_online_avg"], m["starts"],
            m["dangjin_share_pct"], m["net_demand_avg_mw"], m["nuclear_avg_mw"],
            m["coal_target_avg_mw"], m["coal_solved_avg_mw"], m["mismatch_mwh"]]
           + [m["unit_avg_mw"][u] for u in units] for m in summary["monthly"]],
          {1: 10, 3: 16})
    sheet("일별", ["일자", "당진 발전량(MWh)", "당진 평균출력(MW)", "이용률(%)", "평균 운전대수",
                  "기동횟수", "순수요(MW)", "원전(MW)", "석탄목표(MW)", "석탄해(MW)",
                  "목표불일치(MWh)"] + [f"{u}(MWh)" for u in units],
          [[d["date"], d["dangjin_mwh"], d["dangjin_avg_mw"], d["dangjin_cf_pct"],
            d["units_online_avg"], d["starts"], d["net_demand_avg_mw"], d["nuclear_avg_mw"],
            d["coal_target_avg_mw"], d["coal_solved_avg_mw"], d["mismatch_mwh"]]
           + [d["unit_mwh"][u] for u in units] for d in summary["daily"]],
          {1: 12, 2: 16})
    sheet("호기별 월이용률", ["월"] + units,
          [[m["month"]] + [m["unit_cf_pct"][u] for u in units] for m in summary["monthly"]])
    sheet("원전정비", ["호기", "시작", "종료", "비고"],
          [[o["unit"], o["start"], o["end"], o["note"]] for o in summary["outages"]["nuclear"]],
          {4: 40})
    sheet("석탄정비", ["호기", "시작", "종료", "비고"],
          [[o["unit"], o["start"], o["end"], o["note"]] for o in summary["outages"]["coal"]],
          {4: 40})
    info = wb.create_sheet("가정")
    a = summary["annual"]
    rows = [("시나리오", summary["scenario"]["name"]),
            ("기간", f"{summary['scenario']['start']} ~ {summary['scenario']['end']}"),
            ("기상", "평년" if summary["scenario"]["weather"] == "normal"
             else f"{summary['scenario']['weather']}년 실측 재생"),
            ("총수요 증가율(%/년)", round(summary["model_info"]["net_demand"]["demand_growth_pct"], 2)),
            ("태양광 증가율(%/년)", round(summary["model_info"]["net_demand"]["solar_growth_pct"], 2)),
            ("월별 석탄목표 보정(%)", ", ".join(f"{v:g}" for v in summary["scenario"]["coal_adjust_pct"])),
            ("운전예비율", summary["scenario"]["reserve_ratio"]),
            ("순수요 모델 학습", summary["model_info"]["net_demand"]["trained"]),
            ("석탄곡선 모델 학습", summary["model_info"]["coal_curve"]["trained"]),
            ("MILP", json.dumps(summary["model_info"]["milp"], ensure_ascii=False)),
            ("풀이 시간(초)", summary["solve"]["elapsed_seconds"]),
            ("창 상태", json.dumps(summary["solve"]["window_status"], ensure_ascii=False)),
            ("당진 연간 발전량(MWh)", a["dangjin_mwh"]), ("당진 연평균 출력(MW)", a["dangjin_avg_mw"]),
            ("당진 연 이용률(%)", a["dangjin_cf_pct"]), ("당진 연료비(십억원, 2차곡선)", a["dangjin_fuel_cost_billion_won"]),
            ("전국 석탄 해(TWh)", a["coal_solved_twh"]), ("전국 석탄 목표(TWh)", a["coal_target_twh"]),
            ("목표 불일치(GWh)", a["mismatch_gwh"])] + [("경고", w) for w in summary["warnings"]]
    for row in rows:
        info.append(row)
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 90
    wb.move_sheet("가정", offset=-len(wb.sheetnames) + 1)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def save_run(result: AnnualResult, directory: Path) -> dict:
    """runs/<id>/ 에 summary.json, 당진_일월발전량.xlsx, hourly.npz, scenario.json 을 남긴다."""
    directory.mkdir(parents=True, exist_ok=True)
    summary = summarize(result)
    (directory / "summary.json").write_text(json.dumps(summary, ensure_ascii=False), encoding="utf-8")
    result.scenario.save(directory / "scenario.json")
    write_excel(summary, directory / "당진_일월발전량.xlsx")
    np.savez_compressed(directory / "hourly.npz", names=np.asarray(result.names),
                        output=result.output.astype(np.float32), online=result.online,
                        starts=result.starts, demand=result.demand, solar=result.solar,
                        nuclear=result.nuclear, coal_target=result.coal_target,
                        coal_model=result.coal_model, shortage=result.shortage, excess=result.excess)
    return summary


def format_summary(summary: dict) -> str:
    a = summary["annual"]
    lines = [f"[{summary['scenario']['name']}] {summary['scenario']['start']} ~ {summary['scenario']['end']}"
             f"  풀이 {summary['solve']['elapsed_seconds']:.0f}초 · 창 {summary['solve']['window_status']}",
             f"당진본부 연간 {a['dangjin_mwh']/1e3:,.0f} GWh · 연평균 {a['dangjin_avg_mw']:,.0f} MW · "
             f"이용률 {a['dangjin_cf_pct']:.1f}% · 기동 {a['starts']}회",
             "월       당진MWh     평균MW  이용률  운전대수 기동  순수요   원전   석탄목표 석탄해  불일치MWh"]
    for m in summary["monthly"]:
        lines.append(f"{m['month']}  {m['dangjin_mwh']:>10,.0f}  {m['dangjin_avg_mw']:>7,.0f}  "
                     f"{m['dangjin_cf_pct']:5.1f}%  {m['units_online_avg']:5.1f}  {m['starts']:4d}  "
                     f"{m['net_demand_avg_mw']:>6,.0f} {m['nuclear_avg_mw']:>6,.0f}  "
                     f"{m['coal_target_avg_mw']:>6,.0f} {m['coal_solved_avg_mw']:>6,.0f}  {m['mismatch_mwh']:>8,.0f}")
    for w in summary["warnings"]:
        lines.append(f"! {w}")
    return "\n".join(lines)
